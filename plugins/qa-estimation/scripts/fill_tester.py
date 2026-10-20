#!/usr/bin/env python3
"""Fill the "Tester (hrs)" column (and optionally Notes) of the Timedoor quotation sheet.

Usage:
    python3 fill_tester.py QUOTATION.xlsx --show                       # list task rows + current Tester hours
    python3 fill_tester.py QUOTATION.xlsx PLAN.json --backup-dir DIR --dry-run
    python3 fill_tester.py QUOTATION.xlsx PLAN.json --backup-dir DIR

PLAN.json - changes per task row, matched by module code + task text prefix:
    {
      "changes": [
        {"code": "M05", "task": "Checkout page", "hours": 10,
         "expect": 12,                        # optional: current hours; abort if different
         "note": "QA: termasuk ...",          # optional
         "note_mode": "append"}               # "append" (default, joined with " | ") or "replace"
      ]
    }

Why this script exists:
- It writes only Tester/Notes cells via openpyxl. Never re-save this workbook through
  LibreOffice: the Timeline sheet uses LET, which LibreOffice rewrites to let/#NAME?.
- It checks for a LibreOffice lock file. If the user has the file open it still saves but
  warns them to close that window WITHOUT saving. Never kill the user's soffice process.
- It backs up before writing and prints per-section subtotals summed from the cells
  (no recalculation needed). If any change fails to match, nothing is saved.
"""
import argparse
import json
import shutil
import sys
import time
from pathlib import Path

import openpyxl


def find_layout(ws):
    for r in range(1, 20):
        heads = {str(ws.cell(r, c).value).strip(): c for c in range(1, ws.max_column + 1) if ws.cell(r, c).value}
        if any(h.startswith("Tester") for h in heads):
            col = lambda pre: next(c for h, c in heads.items() if h.startswith(pre))
            return r, {"no": 1, "code": col("Code"), "task": col("Task"), "tester": col("Tester"),
                       "notes": next((c for h, c in heads.items() if h.startswith("Notes")), None)}
    sys.exit("Header 'Tester (hrs)' not found on the Quotation sheet")


def task_rows(ws, hr, cols):
    """Yield (row, section, code, task) for task rows up to TOTAL BASE."""
    section = None
    for r in range(hr + 1, ws.max_row + 1):
        a = ws.cell(r, cols["no"]).value
        label = " ".join(str(ws.cell(r, c).value or "") for c in range(1, 4))
        if "TOTAL BASE" in label:
            return
        if a == "▸":
            section = str(ws.cell(r, 3).value or "").strip()
            continue
        code, task = ws.cell(r, cols["code"]).value, ws.cell(r, cols["task"]).value
        if code and task and a is None:
            yield r, section, str(code).strip(), str(task).strip()


def section_totals(ws, hr, cols):
    tot = {}
    for r, sec, _, _ in task_rows(ws, hr, cols):
        tot[sec] = tot.get(sec, 0) + float(ws.cell(r, cols["tester"]).value or 0)
    return tot


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("plan", nargs="?")
    ap.add_argument("--backup-dir")
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--show", action="store_true")
    a = ap.parse_args()

    path = Path(a.xlsx)
    wb = openpyxl.load_workbook(path)
    ws = wb["Quotation"]
    hr, cols = find_layout(ws)

    if a.show:
        for r, sec, code, task in task_rows(ws, hr, cols):
            note = ws.cell(r, cols["notes"]).value if cols["notes"] else None
            print(f"{r:>4} {sec[:14]:<14} {code:<5} T={ws.cell(r, cols['tester']).value!s:<5} {task[:70]}"
                  + (f"  | {note}" if note else ""))
        tot = section_totals(ws, hr, cols)
        print(tot, "TOTAL", sum(tot.values()))
        return

    if not a.plan or not a.backup_dir:
        sys.exit("PLAN.json and --backup-dir are required unless --show is used")

    lock = path.parent / f".~lock.{path.name}#"
    if lock.exists():
        print(f"WARNING: {path.name} is open in LibreOffice. Afterwards, ask the user to close that window "
              "WITHOUT saving and reopen it. Do not kill the soffice process.")

    rows = list(task_rows(ws, hr, cols))
    before = section_totals(ws, hr, cols)
    errors = []
    for ch in json.loads(Path(a.plan).read_text())["changes"]:
        hits = [r for r, _, code, task in rows if code == ch["code"] and task.startswith(ch["task"])]
        if len(hits) != 1:
            errors.append(f"{ch['code']} '{ch['task']}': {len(hits)} rows match")
            continue
        r = hits[0]
        cur = ws.cell(r, cols["tester"]).value
        if "expect" in ch and float(cur or 0) != float(ch["expect"]):
            errors.append(f"{ch['code']} '{ch['task']}': current hours {cur}, plan expects {ch['expect']}")
            continue
        ws.cell(r, cols["tester"]).value = float(ch["hours"])
        if ch.get("note") and cols["notes"]:
            old = ws.cell(r, cols["notes"]).value
            mode = ch.get("note_mode", "append")
            ws.cell(r, cols["notes"]).value = (f"{old} | {ch['note']}" if old and mode == "append" else ch["note"])
        print(f"row {r} {ch['code']} {ch['task'][:40]}: {cur} -> {ch['hours']}")
    if errors:
        sys.exit("Aborted, nothing saved:\n  " + "\n  ".join(errors))

    after = section_totals(ws, hr, cols)
    print("Before:", before, "TOTAL", sum(before.values()))
    print("After: ", after, "TOTAL", sum(after.values()))
    if a.dry_run:
        print("(dry run - file not changed)")
        return
    bdir = Path(a.backup_dir)
    bdir.mkdir(parents=True, exist_ok=True)
    backup = bdir / f"{path.stem}.backup-{time.strftime('%Y%m%d-%H%M%S')}{path.suffix}"
    shutil.copy2(path, backup)
    wb.calculation.fullCalcOnLoad = True
    wb.save(path)
    print(f"Saved. Backup: {backup}")


if __name__ == "__main__":
    main()
