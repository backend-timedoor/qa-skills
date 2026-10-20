#!/usr/bin/env python3
"""Build the companion QA man-hours workbook ("QA Man Hours" + "Asumsi" sheets) from a spec JSON.

Usage:
    python3 build_qa_workbook.py SPEC.json

SPEC.json (full example: ../references/example-spec.json):
{
  "title": "Project name",
  "output": "/abs/path/qa-manhours-<project>.xlsx",
  "params": {"tier1": 0.5, "tier2": 0.25, "buffer": 0.1, "hours_per_day": 8, "qa_count": 1},
  "notes": ["Assumption 1", "Assumption 2"],
  "sections": [
    {"name": "WEB", "modules": [
      {"no": 5, "id": "M05", "name": "Cart & Checkout", "scope": "What is tested ...",
       "tc": 6, "ex": 26, "tier": 1, "note": "Note"},
      {"no": 19, "id": "M19", "name": "Foundation", "scope": "...", "tc": 6, "ex": 25,
       "retest": 4, "note": "Retest = final E2E regression"}
    ]}
  ],
  "optional": {"title": "OPSIONAL - ... (TIDAK termasuk Grand Total)", "modules": [ ... ]}
}

Per module: tc = test-case hours, ex = execution hours (cycle 1). Retest & regression:
  - "retest": fixed number (blue input), or
  - "tier": 1 -> CEILING(ex x tier1, 1), 2 -> CEILING(ex x tier2, 1), 0/absent -> 0.
The quotation's Tester total for a module must equal tc + ex + retest.
The workbook is saved with formulas and fullCalcOnLoad; the script prints the totals it
expects (computed in Python), so no LibreOffice recalculation is needed.
"""
import json
import math
import sys

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

F = "Arial"
NUM = '[=0]"-";General'
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FILL = {k: PatternFill("solid", fgColor=v) for k, v in
        {"hdr": "1F4E78", "sec": "D9E1F2", "sub": "F2F2F2", "tot": "FCE4D6", "gt": "F8CBAD", "in": "FFFF00"}.items()}


def build(spec):
    p = {"tier1": 0.5, "tier2": 0.25, "buffer": 0.1, "hours_per_day": 8, "qa_count": 1, **spec.get("params", {})}
    wb = Workbook()
    ws = wb.active
    ws.title = "QA Man Hours"
    asm = wb.create_sheet("Asumsi")

    # Asumsi: parameter di B5:B9 (dirujuk formula), catatan mulai baris 12
    asm["A1"] = f"Asumsi Perhitungan Man Hours QA (AI-Powered) — {spec['title']}"
    asm["A1"].font = Font(name=F, bold=True, size=13)
    asm["A2"] = "Sel berlatar kuning adalah parameter yang bisa diubah; sheet QA Man Hours menghitung ulang otomatis."
    asm["A2"].font = Font(name=F, italic=True, size=9)
    for c, t in (("A4", "Parameter"), ("B4", "Nilai")):
        asm[c] = t
        asm[c].font, asm[c].fill = Font(name=F, bold=True, color="FFFFFF"), FILL["hdr"]
    params = [("Rasio retest & regression Tier 1 (modul krusial) terhadap eksekusi siklus 1", p["tier1"], "0%"),
              ("Buffer / contingency (internal, tidak masuk quotation)", p["buffer"], "0%"),
              ("Jam kerja per hari (untuk durasi kalender)", p["hours_per_day"], "0"),
              ("Jumlah QA", p["qa_count"], "0"),
              ("Rasio retest & regression Tier 2 (disarankan) terhadap eksekusi siklus 1", p["tier2"], "0%")]
    for i, (k, v, fmt) in enumerate(params, start=5):
        asm.cell(i, 1, k).font = Font(name=F)
        c = asm.cell(i, 2, v)
        c.font, c.fill, c.number_format = Font(name=F, color="0000FF"), FILL["in"], fmt
    asm["A11"] = "Catatan"
    asm["A11"].font = Font(name=F, bold=True)
    for i, n in enumerate(spec.get("notes", []), start=12):
        c = asm.cell(i, 1, f"{i - 11}. {n}")
        c.font, c.alignment = Font(name=F), Alignment(wrap_text=True, vertical="top")
        asm.merge_cells(start_row=i, start_column=1, end_row=i, end_column=2)
        asm.row_dimensions[i].height = 30
    asm.column_dimensions["A"].width, asm.column_dimensions["B"].width = 95, 12
    T1, T2, BUF, HPD, QA = "Asumsi!$B$5", "Asumsi!$B$9", "Asumsi!$B$6", "Asumsi!$B$7", "Asumsi!$B$8"

    ws["A1"] = f"Estimasi Man Hours QA (AI-Powered) — {spec['title']}"
    ws["A1"].font = Font(name=F, bold=True, size=14)
    ws["A2"] = "Angka biru = input estimasi (jam); hitam = formula. Parameter di sheet Asumsi."
    ws["A2"].font = Font(name=F, italic=True, size=9)
    heads = ["No", "ID", "Modul / Aktivitas", "Cakupan Test", "Test Case (Jam)", "Eksekusi Siklus 1 (Jam)",
             "Retest & Regression (Jam)", "Total (Jam)", "Catatan"]
    for j, h in enumerate(heads, 1):
        c = ws.cell(4, j, h)
        c.font, c.fill, c.border = Font(name=F, bold=True, color="FFFFFF"), FILL["hdr"], BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[4].height = 32
    ws["G4"].comment = Comment("Hanya modul krusial: Tier 1 = CEILING(Siklus 1 × rasio Tier 1), Tier 2 = "
                               "CEILING(Siklus 1 × rasio Tier 2). Angka biru = nilai tetap. Modul lain 0.", "QA")
    row = 5

    def band(r, fill):
        for j in range(1, 10):
            ws.cell(r, j).fill, ws.cell(r, j).border = fill, BORDER

    def section(title):
        nonlocal row
        ws.cell(row, 1, "▸").font = Font(name=F, bold=True)
        ws.cell(row, 3, title).font = Font(name=F, bold=True)
        band(row, FILL["sec"])
        row += 1
        return row

    def module(m):
        nonlocal row
        for j, v in enumerate([m.get("no", ""), m["id"], m["name"], m.get("scope", "")], 1):
            c = ws.cell(row, j, v)
            c.font, c.alignment = Font(name=F), Alignment(wrap_text=True, vertical="top")
        ws.cell(row, 5, m["tc"]).font = Font(name=F, color="0000FF")
        ws.cell(row, 6, m["ex"]).font = Font(name=F, color="0000FF")
        if m.get("retest") is not None:
            ws.cell(row, 7, m["retest"]).font = Font(name=F, color="0000FF")
        elif m.get("tier"):
            ws.cell(row, 7, f"=CEILING(F{row}*{T1 if m['tier'] == 1 else T2},1)").font = Font(name=F)
        else:
            ws.cell(row, 7, 0).font = Font(name=F)
        ws.cell(row, 8, f"=SUM(E{row}:G{row})").font = Font(name=F, bold=True)
        c = ws.cell(row, 9, m.get("note", ""))
        c.font, c.alignment = Font(name=F, size=9), Alignment(wrap_text=True, vertical="top")
        for j in range(1, 10):
            ws.cell(row, j).border = BORDER
        for j in range(5, 9):
            ws.cell(row, j).number_format = NUM
            ws.cell(row, j).alignment = Alignment(horizontal="center", vertical="top")
        row += 1

    def subtotal(label, first):
        nonlocal row
        ws.cell(row, 1, "Σ").font = Font(name=F, bold=True)
        ws.cell(row, 3, label).font = Font(name=F, bold=True)
        for col in "EFGH":
            c = ws[f"{col}{row}"]
            c.value = f"=SUM({col}{first}:{col}{row - 1})"
            c.font, c.number_format, c.alignment = Font(name=F, bold=True), NUM, Alignment(horizontal="center")
        band(row, FILL["sub"])
        row += 1
        return row - 1

    def total(label, formula, fill, bold=True, fmt="General", note=None):
        nonlocal row
        ws.cell(row, 3, label).font = Font(name=F, bold=bold)
        c = ws.cell(row, 8, formula)
        c.font, c.number_format, c.alignment = Font(name=F, bold=bold), fmt, Alignment(horizontal="center")
        if note:
            ws.cell(row, 9, note).font = Font(name=F, size=9)
        band(row, fill)
        row += 1
        return row - 1

    subs = []
    for sec in spec["sections"]:
        first = section(sec["name"])
        for m in sec["modules"]:
            module(m)
        subs.append(subtotal(f"Subtotal {sec['name']}", first))
    row += 1
    base = total("Total sebelum buffer", "=" + "+".join(f"H{r}" for r in subs), FILL["tot"],
                 note="Sama dengan total kolom Tester di quotation")
    buf = total("Buffer / contingency (internal)", f"=CEILING(H{base}*{BUF},1)", FILL["tot"], bold=False,
                note="Total × buffer (sheet Asumsi); tidak masuk quotation")
    gt = total("GRAND TOTAL QA (Jam)", f"=H{base}+H{buf}", FILL["gt"])
    total("Estimasi durasi kalender (minggu, 5 hari kerja)", f"=H{gt}/{HPD}/{QA}/5", FILL["tot"], bold=False,
          fmt="0.0", note="Grand Total ÷ jam per hari ÷ jumlah QA ÷ 5")

    opt = spec.get("optional")
    if opt and opt.get("modules"):
        row += 2
        first = section(opt["title"])
        for m in opt["modules"]:
            module(m)
        o = subtotal("Subtotal OPSIONAL (sebelum buffer)", first)
        ob = total("Subtotal OPSIONAL + buffer", f"=H{o}+CEILING(H{o}*{BUF},1)", FILL["tot"])
        total("GRAND TOTAL jika opsional diambil (Jam)", f"=H{gt}+H{ob}", FILL["gt"])

    for k, v in {"A": 5, "B": 8, "C": 34, "D": 62, "E": 11, "F": 12, "G": 13, "H": 11, "I": 34}.items():
        ws.column_dimensions[k].width = v
    ws.freeze_panes = "A5"
    ws.page_setup.orientation = "landscape"
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    wb.calculation.fullCalcOnLoad = True
    wb.save(spec["output"])
    print(spec["output"])

    def mod_total(m):
        rt = m.get("retest")
        if rt is None:
            rate = {1: p["tier1"], 2: p["tier2"]}.get(m.get("tier") or 0, 0)
            rt = math.ceil(m["ex"] * rate - 1e-9) if rate else 0
        return m["tc"] + m["ex"] + rt
    secs = {s["name"]: sum(mod_total(m) for m in s["modules"]) for s in spec["sections"]}
    base_total = sum(secs.values())
    gt_total = base_total + math.ceil(base_total * p["buffer"] - 1e-9)
    print("Subtotals:", secs)
    print(f"Total before buffer (= quotation Tester total): {base_total}")
    print(f"Grand total incl. internal buffer: {gt_total} (~{gt_total / p['hours_per_day'] / p['qa_count'] / 5:.1f} weeks)")
    if opt and opt.get("modules"):
        print(f"Optional subtotal: {sum(mod_total(m) for m in opt['modules'])}")


if __name__ == "__main__":
    build(json.loads(open(sys.argv[1]).read()))
